# -*- coding: utf-8 -*-
"""
Construit le classeur (feuilles, styles, tableaux, formules, validations,
mise en forme conditionnelle, mise en page, protection, plages nommees)
avec openpyxl, puis l'enregistre en .xlsx.

L'assemblage final en .xlsm (injection du vbaProject.bin, types de contenu,
relations, codeName) est realise par assemble.py.
"""
import datetime

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, Protection
)
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

import sp_config as C

# ---------------------------------------------------------------------------
# Raccourcis de style
# ---------------------------------------------------------------------------
def _f(**kw):
    kw.setdefault("name", C.FONT_NAME)
    return Font(**kw)

def _fill(hexcolor):
    return PatternFill(start_color=hexcolor, end_color=hexcolor, fill_type="solid")

def _side(color=C.CLR_BORDER, style="thin"):
    return Side(style=style, color=color)

def _border(color=C.CLR_BORDER, style="thin"):
    s = _side(color, style)
    return Border(left=s, right=s, top=s, bottom=s)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
UNLOCKED = Protection(locked=False)


def _title_band(ws, cell_range, text, size=18, bg=C.CLR_PRIMARY, fg="FFFFFF"):
    """Bandeau de titre fusionne."""
    ws.merge_cells(cell_range)
    top_left = cell_range.split(":")[0]
    c = ws[top_left]
    c.value = text
    c.font = _f(size=size, bold=True, color=fg)
    c.alignment = CENTER
    c.fill = _fill(bg)
    # applique le fond a toutes les cellules fusionnees
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = _fill(bg)


def _table_header(ws, row, first_col, headers, bg=C.CLR_HEADER_BG, fg="FFFFFF"):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=first_col + i, value=h)
        c.font = _f(bold=True, color=fg, size=11)
        c.alignment = CENTER
        c.fill = _fill(bg)
        c.border = _border(C.CLR_HEADER_BG)


# ===========================================================================
def build(xlsx_path):
    wb = Workbook()
    # supprime la feuille par defaut
    wb.remove(wb.active)

    # cree les feuilles dans l'ordre et fixe le codeName VBA
    sheets = {}
    for tab, cn in C.SHEETS:
        ws = wb.create_sheet(title=tab)
        ws.sheet_properties.codeName = cn
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = C.CLR_PRIMARY
        sheets[cn] = ws

    today = datetime.date.today()

    # donnees d'exemple
    employees = [
        ("E001", "BENANI Youssef", "Chef de chantier", "Équipe A", "Actif"),
        ("E002", "EL AMRANI Karim", "Maçon", "Équipe A", "Actif"),
        ("E003", "TAZI Rachid", "Manœuvre", "Équipe A", "Actif"),
        ("E004", "OUALI Hamid", "Grutier", "Équipe B", "Actif"),
        ("E005", "BENNANI Saïd", "Électricien", "Équipe B", "Actif"),
        ("E006", "CHRAIBI Omar", "Plombier", "Équipe B", "Actif"),
        ("E007", "IDRISSI Mohamed", "Charpentier", "Équipe A", "Actif"),
        ("E008", "ALAOUI Nabil", "Peintre", "Équipe B", "Actif"),
        ("E009", "FASSI Brahim", "Manœuvre", "Équipe A", "Inactif"),
    ]
    actifs = [e for e in employees if e[4] == "Actif"]
    present_ids = {e[0] for e in actifs[:6]}   # 6 presents pour la demo

    _build_employes(sheets[C.CN_EMP], employees)
    _build_presence(sheets[C.CN_PRES], actifs, today)
    _build_impression(sheets[C.CN_IMP], actifs, present_ids, today)
    _build_historique(sheets[C.CN_HIST], actifs, present_ids, today)
    _build_parametres(sheets[C.CN_PARAM])
    _build_accueil(sheets[C.CN_ACCUEIL])
    _build_guide(sheets[C.CN_GUIDE])

    _defined_names(wb)

    # feuille active a l'ouverture
    wb.active = wb.sheetnames.index(C.TAB_ACCUEIL)

    wb.save(xlsx_path)
    return xlsx_path


# ---------------------------------------------------------------------------
def _build_employes(ws, employees):
    E = C.Emp
    ws.column_dimensions["A"].width = 2
    widths = {"B": 14, "C": 30, "D": 20, "E": 14, "F": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    _title_band(ws, "B2:F2", "LISTE DES EMPLOYÉS")
    ws.row_dimensions[2].height = 30

    note = ws.cell(row=3, column=2,
                   value="Ajoutez, modifiez ou désactivez vos employés ici. "
                         "Les changements se répercutent automatiquement dans tout le classeur.")
    note.font = _f(italic=True, color=C.CLR_MUTED, size=10)
    ws.merge_cells("B3:F3")
    ws.row_dimensions[3].height = 22

    _table_header(ws, E.HEADER_ROW, E.COL_ID, E.HEADERS)

    r = E.FIRST_ROW
    for (eid, nom, poste, equipe, statut) in employees:
        ws.cell(row=r, column=E.COL_ID, value=eid).alignment = CENTER
        ws.cell(row=r, column=E.COL_NOM, value=nom).alignment = LEFT
        ws.cell(row=r, column=E.COL_POSTE, value=poste).alignment = LEFT
        ws.cell(row=r, column=E.COL_EQUIPE, value=equipe).alignment = CENTER
        ws.cell(row=r, column=E.COL_STATUT, value=statut).alignment = CENTER
        r += 1
    last = r - 1

    ref = f"{get_column_letter(E.COL_ID)}{E.HEADER_ROW}:{get_column_letter(E.COL_STATUT)}{last}"
    tbl = Table(displayName=E.TABLE, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False,
        showFirstColumn=False, showLastColumn=False)
    ws.add_table(tbl)

    # validation Statut (Actif/Inactif)
    dv = DataValidation(type="list", formula1='"Actif,Inactif"', allow_blank=False,
                        showErrorMessage=True)
    dv.error = "Choisissez Actif ou Inactif."
    dv.prompt = "Actif ou Inactif"
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(E.COL_STATUT)}{E.FIRST_ROW}:"
           f"{get_column_letter(E.COL_STATUT)}1000")

    ws.freeze_panes = f"A{E.HEADER_ROW + 1}"
    # feuille Employes : editable (pas de protection)


# ---------------------------------------------------------------------------
def _build_presence(ws, actifs, today):
    P = C.Pres
    ws.column_dimensions["A"].width = 2
    for col, w in {"B": 12, "C": 16, "D": 32, "E": 22}.items():
        ws.column_dimensions[col].width = w

    _title_band(ws, "B2:E2", "PRÉSENCE DU JOUR")
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[P.BTN_ROW].height = 42   # bandeau des boutons (VBA)

    # Date / Chantier
    ld = ws[P.DATE_LBL]; ld.value = "Date :"; ld.font = _f(bold=True); ld.alignment = RIGHT
    cd = ws[P.DATE_CELL]; cd.value = today
    cd.number_format = "dd/mm/yyyy"; cd.font = _f(bold=True, size=12); cd.alignment = CENTER
    cd.fill = _fill(C.CLR_CARD_BG); cd.border = _border(); cd.protection = UNLOCKED

    lc = ws[P.CHANTIER_LBL]; lc.value = "Chantier :"; lc.font = _f(bold=True); lc.alignment = RIGHT
    cc = ws[P.CHANTIER_CELL]; cc.value = "=param_Chantier"
    cc.font = _f(size=12); cc.alignment = LEFT
    cc.fill = _fill(C.CLR_CARD_BG); cc.border = _border(); cc.protection = UNLOCKED

    # Recherche
    ls = ws[P.SEARCH_LBL]; ls.value = "Rechercher :"; ls.font = _f(bold=True); ls.alignment = RIGHT
    cs = ws[P.SEARCH_CELL]; cs.value = None
    cs.fill = _fill("FFFFFF"); cs.border = _border(); cs.protection = UNLOCKED
    cs.alignment = LEFT

    # Cartes KPI
    def kpi(lbl_cell, val_cell, label, formula, color):
        lc = ws[lbl_cell]; lc.value = label
        lc.font = _f(bold=True, color="FFFFFF", size=9); lc.alignment = CENTER
        lc.fill = _fill(color)
        vc = ws[val_cell]; vc.value = formula
        vc.font = _f(bold=True, size=20, color=color); vc.alignment = CENTER
        vc.fill = _fill(C.CLR_CARD_BG); vc.border = _border(color)

    rng = f"$B${P.FIRST_ROW}:$B${P.MAX_ROW}"
    kpi(P.KPI_ACTIF_LBL, P.KPI_ACTIF_VAL, "EMPLOYÉS ACTIFS",
        '=COUNTIF(tblEmployes[Statut],"Actif")', C.CLR_PRIMARY)
    kpi(P.KPI_PRES_LBL, P.KPI_PRES_VAL, "PRÉSENTS",
        f"=COUNTIF({rng},UNICHAR({C.CHECK_CODE}))", C.CLR_ACCENT)
    kpi(P.KPI_ABS_LBL, P.KPI_ABS_VAL, "ABSENTS",
        f"={P.KPI_ACTIF_VAL}-{P.KPI_PRES_VAL}", C.CLR_RED)
    ws.row_dimensions[P.KPI_LBL_ROW].height = 16
    ws.row_dimensions[P.KPI_VAL_ROW].height = 30

    # En-tete de la liste
    _table_header(ws, P.HEADER_ROW, P.COL_CHK, P.HEADERS)
    ws.row_dimensions[P.HEADER_ROW].height = 22

    # Pre-remplissage (visible meme sans macro) : employes actifs, non coches
    r = P.FIRST_ROW
    for (eid, nom, poste, equipe, statut) in actifs:
        ws.cell(row=r, column=P.COL_CHK, value=None).alignment = CENTER
        ws.cell(row=r, column=P.COL_ID, value=eid).alignment = CENTER
        ws.cell(row=r, column=P.COL_NOM, value=nom).alignment = LEFT
        ws.cell(row=r, column=P.COL_POSTE, value=poste).alignment = LEFT
        r += 1

    # Mise en forme conditionnelle sur toute la zone
    first, last = P.FIRST_ROW, P.MAX_ROW
    cchk = get_column_letter(P.COL_CHK)
    cid = get_column_letter(P.COL_ID)
    zone = (f"{cchk}{first}:{get_column_letter(P.COL_POSTE)}{last}")

    # 1) bordures sur les lignes remplies (priorite haute, ne stoppe pas)
    ws.conditional_formatting.add(zone, FormulaRule(
        formula=[f'${cid}{first}<>""'],
        border=_border(C.CLR_BORDER), stopIfTrue=False))
    # 2) surlignage vert des presents (stoppe -> pas de rayure par-dessus)
    ws.conditional_formatting.add(zone, FormulaRule(
        formula=[f'${cchk}{first}=UNICHAR({C.CHECK_CODE})'],
        fill=_fill(C.CLR_PRESENT_BG), font=_f(bold=True), stopIfTrue=True))
    # 3) lignes alternees
    ws.conditional_formatting.add(zone, FormulaRule(
        formula=[f'AND(${cid}{first}<>"",MOD(ROW(),2)=0)'],
        fill=_fill(C.CLR_BAND), stopIfTrue=False))

    # Coche centree et grande sur la colonne Present
    for rr in range(first, last + 1):
        cc = ws.cell(row=rr, column=P.COL_CHK)
        cc.alignment = CENTER
        cc.font = _f(size=13, bold=True, color=C.CLR_ACCENT_DK)

    ws.freeze_panes = P.FREEZE

    # Protection : seules les cellules de saisie sont deverrouillees
    ws.protection.sheet = True
    ws.protection.password = C.PROTECT_PW
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.protection.autoFilter = False


# ---------------------------------------------------------------------------
def _build_impression(ws, actifs, present_ids, today):
    I = C.Imp
    ws.column_dimensions["A"].width = 2
    for col, w in {"B": 8, "C": 16, "D": 34, "E": 24, "F": 3, "G": 22}.items():
        ws.column_dimensions[col].width = w

    _title_band(ws, "B1:E1", "=param_Entreprise", size=16, bg="FFFFFF", fg=C.CLR_PRIMARY)
    ws.row_dimensions[1].height = 26
    _title_band(ws, "B2:E2", "=param_TitreImpression", size=14, bg=C.CLR_PRIMARY, fg="FFFFFF")
    ws.row_dimensions[2].height = 26

    # Date (selecteur) / Chantier
    ws[I.DATE_LBL].value = "Date :"; ws[I.DATE_LBL].font = _f(bold=True); ws[I.DATE_LBL].alignment = RIGHT
    cd = ws[I.DATE_CELL]; cd.value = today
    cd.number_format = "dd/mm/yyyy"; cd.font = _f(bold=True, size=12); cd.alignment = CENTER
    cd.fill = _fill(C.CLR_CARD_BG); cd.border = _border(); cd.protection = UNLOCKED

    ws[I.CHANTIER_LBL].value = "Chantier :"; ws[I.CHANTIER_LBL].font = _f(bold=True); ws[I.CHANTIER_LBL].alignment = RIGHT
    cc = ws[I.CHANTIER_CELL]; cc.value = "=param_Chantier"
    cc.font = _f(size=12); cc.alignment = LEFT

    # En-tete du tableau imprimable
    _table_header(ws, I.HEADER_ROW, I.COL_NUM, I.HEADERS)
    ws.row_dimensions[I.HEADER_ROW].height = 22

    # Pre-remplissage demo : presents du jour, tries par nom
    present = sorted([e for e in actifs if e[0] in present_ids], key=lambda e: e[1])
    r = I.FIRST_ROW
    for i, (eid, nom, poste, equipe, statut) in enumerate(present, start=1):
        ws.cell(row=r, column=I.COL_NUM, value=i).alignment = CENTER
        ws.cell(row=r, column=I.COL_ID, value=eid).alignment = CENTER
        ws.cell(row=r, column=I.COL_NOM, value=nom).alignment = LEFT
        ws.cell(row=r, column=I.COL_POSTE, value=poste).alignment = LEFT
        for cc2 in range(I.COL_NUM, I.COL_POSTE + 1):
            ws.cell(row=r, column=cc2).border = _border("969FA8")
        if i % 2 == 0:
            for cc2 in range(I.COL_NUM, I.COL_POSTE + 1):
                ws.cell(row=r, column=cc2).fill = _fill(C.CLR_BAND)
        r += 1
    tot_row = r
    ws.cell(row=tot_row, column=I.COL_ID, value="Total présents :").font = _f(bold=True)
    ws.cell(row=tot_row, column=I.COL_ID).alignment = RIGHT
    tc = ws.cell(row=tot_row, column=I.COL_POSTE, value=len(present))
    tc.font = _f(bold=True); tc.alignment = CENTER

    # colonne masquee des dates uniques (pour la liste deroulante)
    hc = get_column_letter(I.HELPER_COL)
    ws[f"{hc}1"].value = "Dates"
    ws[f"{hc}2"].value = f"=IFERROR(SORT(UNIQUE({C.Hist.TABLE}[Date]),1,-1),\"\")"
    ws.column_dimensions[hc].hidden = True
    for rr in range(2, 400):
        ws.cell(row=rr, column=I.HELPER_COL).number_format = "dd/mm/yyyy"

    # validation : liste deroulante des dates enregistrees (plage dynamique)
    dv = DataValidation(type="list", formula1=f"${hc}$2#", allow_blank=True,
                        showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(I.DATE_CELL)

    # Mise en page A4 portrait, ajustee sur une page en largeur
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0   # largeur sur 1 page ; hauteur = autant que necessaire
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.print_area = f"B1:E{tot_row + 1}"

    # Protection : seul le selecteur de date est deverrouille
    ws.protection.sheet = True
    ws.protection.password = C.PROTECT_PW


# ---------------------------------------------------------------------------
def _build_historique(ws, actifs, present_ids, today):
    H = C.Hist
    ws.column_dimensions["A"].width = 2
    for col, w in {"B": 13, "C": 12, "D": 28, "E": 20, "F": 20, "G": 10}.items():
        ws.column_dimensions[col].width = w

    _title_band(ws, "B2:G2", "HISTORIQUE DES PRÉSENCES")
    ws.row_dimensions[2].height = 30
    note = ws.cell(row=3, column=2,
                   value="Alimenté automatiquement lors de l'enregistrement. "
                         "Ne pas modifier manuellement. L'historique n'est jamais supprimé automatiquement.")
    note.font = _f(italic=True, color=C.CLR_MUTED, size=10)
    ws.merge_cells("B3:G3")
    ws.row_dimensions[3].height = 22

    _table_header(ws, H.HEADER_ROW, H.COL_DATE, H.HEADERS)

    # Donnees demo : une journee (aujourd'hui a la construction du fichier)
    r = H.FIRST_ROW
    default_chantier = C.Param.DEF_CHANTIER
    for (eid, nom, poste, equipe, statut) in actifs:
        pres = "Oui" if eid in present_ids else "Non"
        dc = ws.cell(row=r, column=H.COL_DATE, value=today); dc.number_format = "dd/mm/yyyy"
        dc.alignment = CENTER
        ws.cell(row=r, column=H.COL_ID, value=eid).alignment = CENTER
        ws.cell(row=r, column=H.COL_NOM, value=nom).alignment = LEFT
        ws.cell(row=r, column=H.COL_POSTE, value=poste).alignment = LEFT
        ws.cell(row=r, column=H.COL_CHANTIER, value=default_chantier).alignment = LEFT
        ws.cell(row=r, column=H.COL_PRESENT, value=pres).alignment = CENTER
        r += 1
    last = r - 1

    ref = f"{get_column_letter(H.COL_DATE)}{H.HEADER_ROW}:{get_column_letter(H.COL_PRESENT)}{last}"
    tbl = Table(displayName=H.TABLE, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tbl)

    ws.freeze_panes = f"A{H.HEADER_ROW + 1}"
    ws.protection.sheet = True
    ws.protection.password = C.PROTECT_PW


# ---------------------------------------------------------------------------
def _build_parametres(ws):
    Pm = C.Param
    ws.column_dimensions["A"].width = 2
    for col, w in {"B": 34, "C": 22, "D": 14, "E": 14}.items():
        ws.column_dimensions[col].width = w

    _title_band(ws, "B2:E2", "PARAMÈTRES DU CLASSEUR")
    ws.row_dimensions[2].height = 30

    rows = [
        (Pm.ROW_ENT, "Nom de l'entreprise", Pm.DEF_ENT, None),
        (Pm.ROW_CHANTIER, "Nom du chantier par défaut", Pm.DEF_CHANTIER, None),
        (Pm.ROW_TITRE, "Titre du document imprimé", Pm.DEF_TITRE, None),
        (Pm.ROW_PDF, "Dossier d'export PDF (vide = dossier du fichier)", Pm.DEF_PDF, None),
        (Pm.ROW_ORIENT, "Orientation d'impression", Pm.DEF_ORIENT, "orient"),
        (Pm.ROW_POINTAGE, "Méthode de pointage", "Double-clic sur la ligne (ou boutons Tout cocher / décocher)", "locked"),
    ]
    for (r, label, default, kind) in rows:
        lc = ws.cell(row=r, column=Pm.LBL_COL, value=label)
        lc.font = _f(bold=True, color=C.CLR_TEXT); lc.alignment = LEFT
        lc.fill = _fill(C.CLR_CARD_BG)
        ws.merge_cells(start_row=r, start_column=Pm.VAL_COL, end_row=r, end_column=5)
        vc = ws.cell(row=r, column=Pm.VAL_COL, value=default)
        vc.alignment = LEFT
        vc.border = _border()
        if kind == "locked":
            vc.font = _f(italic=True, color=C.CLR_MUTED)
            vc.fill = _fill(C.CLR_CARD_BG)
        else:
            vc.font = _f()
            vc.fill = _fill("FFFFFF")
            vc.protection = UNLOCKED
        ws.row_dimensions[r].height = 20

    # validation orientation
    dv = DataValidation(type="list", formula1='"Portrait,Paysage"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(Pm.VAL_COL)}{Pm.ROW_ORIENT}")

    # Zone logo
    lbl = ws.cell(row=Pm.LOGO_ROW - 1, column=Pm.LBL_COL, value="Logo de l'entreprise :")
    lbl.font = _f(bold=True, color=C.CLR_TEXT)
    ws.merge_cells(start_row=Pm.LOGO_ROW, start_column=Pm.LBL_COL,
                   end_row=Pm.LOGO_ROW + 3, end_column=5)
    logo = ws.cell(row=Pm.LOGO_ROW, column=Pm.LBL_COL,
                   value="Insérez votre logo ici  (menu Insertion > Images)")
    logo.alignment = CENTER
    logo.font = _f(italic=True, color=C.CLR_MUTED)
    for row in ws[f"B{Pm.LOGO_ROW}:E{Pm.LOGO_ROW + 3}"]:
        for cell in row:
            cell.border = _border()

    ws.protection.sheet = True
    ws.protection.password = C.PROTECT_PW


# ---------------------------------------------------------------------------
def _build_accueil(ws):
    for col, w in {"A": 2, "B": 4, "C": 20, "D": 16, "E": 4, "F": 6, "G": 18, "H": 10}.items():
        ws.column_dimensions[col].width = w

    _title_band(ws, "B2:H2", "SYSTÈME DE GESTION DES PRÉSENCES", size=20)
    ws.row_dimensions[2].height = 40
    _title_band(ws, "B3:H3", "=param_Entreprise", size=13, bg=C.CLR_PRIMARY_LIGHT, fg="FFFFFF")
    ws.row_dimensions[3].height = 22

    a = ws.cell(row=5, column=2, value="ACTIONS RAPIDES")
    a.font = _f(bold=True, size=12, color=C.CLR_PRIMARY)
    n = ws.cell(row=5, column=6, value="NAVIGATION")
    n.font = _f(bold=True, size=12, color=C.CLR_PRIMARY)

    # panneau clair derriere les boutons (les boutons sont dessines par VBA)
    for row in ws["B6:I17"]:
        for cell in row:
            cell.fill = _fill(C.CLR_CARD_BG)
    for r in range(6, 18):
        ws.row_dimensions[r].height = 18

    # petites instructions
    steps = ws.cell(row=19, column=2,
        value="Chaque matin : 1) Nouvelle journée  2) Double-cliquez les présents  "
              "3) Enregistrer la journée  4) Imprimer.")
    steps.font = _f(italic=True, color=C.CLR_MUTED, size=10)
    ws.merge_cells("B19:H19")
    note = ws.cell(row=20, column=2,
        value="Astuce : si les boutons n'apparaissent pas, activez les macros "
              "(bandeau jaune en haut > Activer le contenu).")
    note.font = _f(italic=True, color=C.CLR_MUTED, size=9)
    ws.merge_cells("B20:H20")

    ws.protection.sheet = True
    ws.protection.password = C.PROTECT_PW


# ---------------------------------------------------------------------------
def _build_guide(ws):
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 3
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 15

    _title_band(ws, "B2:H2", "GUIDE D'UTILISATION", size=18)
    ws.row_dimensions[2].height = 30

    sections = [
        ("h", "1. PREMIÈRE OUVERTURE  (à faire une seule fois)"),
        ("p", "• À l'ouverture, un bandeau jaune peut apparaître : cliquez sur « Activer le contenu » pour autoriser les macros. C'est indispensable pour que les boutons fonctionnent."),
        ("p", "• Enregistrez le fichier au format .xlsm (Excel avec prise en charge des macros) : Fichier > Enregistrer sous > type « Classeur Excel (prenant en charge les macros) »."),
        ("h", "2. CONFIGURER LES EMPLOYÉS"),
        ("p", "• Ouvrez la feuille « Employés »."),
        ("p", "• Pour AJOUTER un employé : écrivez directement sur la première ligne vide sous le tableau. Renseignez l'ID (déjà existant chez vous), le Nom et Prénom, le Poste, l'Équipe et le Statut."),
        ("p", "• Pour MODIFIER un employé : changez simplement le texte de la cellule. La modification apparaît partout automatiquement."),
        ("p", "• Pour DÉSACTIVER un employé : mettez son Statut sur « Inactif ». Il n'apparaîtra plus dans la présence du jour, mais reste dans l'historique."),
        ("h", "3. PARAMÈTRES"),
        ("p", "• Feuille « Paramètres » : renseignez le nom de l'entreprise, le chantier par défaut, le titre imprimé et éventuellement un dossier d'export PDF. Vous pouvez aussi y insérer votre logo."),
        ("h", "4. UTILISATION QUOTIDIENNE  (moins de 2 minutes)"),
        ("p", "1) Ouvrez le fichier."),
        ("p", "2) Cliquez sur « Nouvelle journée » (met la date du jour et décoche tout)."),
        ("p", "3) Sur « Présence du Jour », DOUBLE-CLIQUEZ sur la ligne de chaque employé présent (une coche verte apparaît). Vous pouvez aussi utiliser « Tout cocher » / « Tout décocher »."),
        ("p", "4) Cliquez sur « Enregistrer la journée » (sauvegarde dans l'historique, sans doublon)."),
        ("p", "5) Cliquez sur « Imprimer » (la liste des présents s'affiche et l'impression démarre) ou « Exporter en PDF »."),
        ("h", "5. RECHERCHE ET COMPTEURS"),
        ("p", "• Tapez un nom dans la case « Rechercher » pour filtrer la liste. Effacez la case pour tout réafficher."),
        ("p", "• Les compteurs en haut affichent en temps réel le nombre d'employés actifs, de présents et d'absents."),
        ("h", "6. CONSULTER OU RÉIMPRIMER UNE ANCIENNE JOURNÉE"),
        ("p", "• Ouvrez la feuille « Liste à imprimer »."),
        ("p", "• Choisissez une date dans le menu déroulant (à côté de « Date »). La liste se met à jour automatiquement avec les présents de ce jour-là."),
        ("p", "• Cliquez sur « Imprimer cette liste » ou « Exporter en PDF ». Changer la date ne modifie jamais les données enregistrées."),
        ("h", "7. SAUVEGARDE ET BONNES PRATIQUES"),
        ("p", "• Le fichier est sauvegardé automatiquement avant chaque impression."),
        ("p", "• Enregistrez régulièrement (Ctrl+S) et conservez une copie de sauvegarde du fichier .xlsm (clé USB, cloud, etc.)."),
        ("p", "• Ne supprimez pas la feuille « Historique » : elle contient toutes vos journées passées."),
    ]

    r = 4
    for kind, text in sections:
        cell = ws.cell(row=r, column=2, value=text)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        if kind == "h":
            cell.font = _f(bold=True, size=12, color="FFFFFF")
            for cc in range(2, 9):
                ws.cell(row=r, column=cc).fill = _fill(C.CLR_PRIMARY_LIGHT)
            cell.alignment = LEFT
            ws.row_dimensions[r].height = 22
        else:
            cell.font = _f(size=11, color=C.CLR_TEXT)
            cell.alignment = LEFT_TOP
            ws.row_dimensions[r].height = 30
        r += 1

    ws.protection.sheet = True
    ws.protection.password = C.PROTECT_PW


# ---------------------------------------------------------------------------
def _defined_names(wb):
    def q(tab):
        return "'" + tab + "'"

    names = {
        C.Param.NAME_ENT: f"{q(C.TAB_PARAM)}!$C${C.Param.ROW_ENT}",
        C.Param.NAME_CHANTIER: f"{q(C.TAB_PARAM)}!$C${C.Param.ROW_CHANTIER}",
        C.Param.NAME_TITRE: f"{q(C.TAB_PARAM)}!$C${C.Param.ROW_TITRE}",
        C.Param.NAME_PDF: f"{q(C.TAB_PARAM)}!$C${C.Param.ROW_PDF}",
        C.Param.NAME_ORIENT: f"{q(C.TAB_PARAM)}!$C${C.Param.ROW_ORIENT}",
        C.NAME_DATE_PRES: f"{q(C.TAB_PRES)}!${C.Pres.DATE_CELL[0]}${C.Pres.DATE_CELL[1:]}",
        C.NAME_CHANTIER_PRES: f"{q(C.TAB_PRES)}!${C.Pres.CHANTIER_CELL[0]}${C.Pres.CHANTIER_CELL[1:]}",
        C.NAME_SEARCH_PRES: f"{q(C.TAB_PRES)}!${C.Pres.SEARCH_CELL[0]}${C.Pres.SEARCH_CELL[1:]}",
        C.NAME_DATE_IMP: f"{q(C.TAB_IMP)}!${C.Imp.DATE_CELL[0]}${C.Imp.DATE_CELL[1:]}",
    }
    for nm, ref in names.items():
        wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


if __name__ == "__main__":
    import os
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_tmp_build.xlsx")
    build(out)
    print("xlsx ->", out, os.path.getsize(out))
