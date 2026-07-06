# -*- coding: utf-8 -*-
"""
Orchestrateur : construit le classeur .xlsm complet et le valide.

Usage :
    python build.py [chemin_de_sortie.xlsm]

Doit etre execute avec Python 3.12+ (les librairies ms_ovba utilisent
la syntaxe d'annotations 3.10+).
"""
import os
import sys
import tempfile

import sp_config as C
import make_vba
import workbook
import assemble

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)


def build(output_path=None):
    output_path = output_path or os.path.join(REPO, C.OUTPUT_XLSM)
    workdir = tempfile.mkdtemp(prefix="presence_build_")
    xlsx_path = os.path.join(workdir, "workbook.xlsx")
    vba_path = os.path.join(workdir, "vbaProject.bin")

    print("1/3  Construction du vbaProject.bin ...")
    make_vba.build_vba_bin(vba_path)
    print("       ->", os.path.getsize(vba_path), "octets")

    print("2/3  Construction du classeur (openpyxl) ...")
    workbook.build(xlsx_path)
    print("       ->", os.path.getsize(xlsx_path), "octets")

    print("3/3  Assemblage du .xlsm ...")
    assemble.assemble(xlsx_path, vba_path, output_path)
    print("       ->", output_path, os.path.getsize(output_path), "octets")

    return output_path


def validate(xlsm_path):
    print("\nVALIDATION")
    ok = True

    # 1) ZIP integre + parties essentielles
    import zipfile
    with zipfile.ZipFile(xlsm_path) as z:
        bad = z.testzip()
        assert bad is None, f"Entree ZIP corrompue : {bad}"
        names = z.namelist()
        for req in ["xl/vbaProject.bin", "xl/workbook.xml",
                    "[Content_Types].xml", "xl/_rels/workbook.xml.rels"]:
            assert req in names, f"Partie manquante : {req}"
        ct = z.read("[Content_Types].xml").decode("utf-8")
        assert "macroEnabled.main+xml" in ct, "Type de classeur non macro-enabled"
        assert "vbaProject" in ct, "Type vbaProject manquant"
        wbxml = z.read("xl/workbook.xml").decode("utf-8")
        assert 'codeName="ThisWorkbook"' in wbxml, "codeName ThisWorkbook manquant"
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        assert "vbaProject.bin" in rels, "Relation vbaProject manquante"
    print("  [OK] Structure ZIP / Content_Types / rels / codeName")

    # 2) openpyxl peut rouvrir en conservant le VBA
    from openpyxl import load_workbook
    wb = load_workbook(xlsm_path, keep_vba=True)
    assert wb.vba_archive is not None, "vba_archive absent apres relecture"
    assert set(wb.sheetnames) == {t for t, _ in C.SHEETS}, "Feuilles inattendues"
    print("  [OK] Relecture openpyxl (keep_vba) - feuilles :", wb.sheetnames)

    # 3) olevba : les macros sont extractibles
    try:
        from oletools.olevba import VBA_Parser
        vp = VBA_Parser(xlsm_path)
        mods = []
        for (_, _, name, _) in vp.extract_macros():
            for ext in (".cls", ".bas", ".frm"):
                if name.lower().endswith(ext):
                    name = name[: -len(ext)]
            mods.append(name)
        vp.close()
        need = {"ThisWorkbook", "wsPresence", "modActions", "modUI",
                "modData", "modConfig", "modUtil"}
        missing = need - set(mods)
        assert not missing, f"Modules VBA manquants : {missing}"
        print("  [OK] olevba - modules :", sorted(mods))
    except ImportError:
        print("  [!!] oletools indisponible, controle VBA ignore")

    print("\nRESULTAT : classeur valide ->", xlsm_path)
    return ok


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else None
    path = build(out)
    validate(path)
