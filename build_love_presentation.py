"""
Generate an elegant, professional 5-slide presentation about "Love".

Design language
---------------
- 16:9 widescreen (13.333 x 7.5 in)
- Soft palette: white, blush pink, light red/rose, gold accents
- Serif display headings (Georgia) + clean sans body (Segoe UI)
- Minimal text, generous whitespace, consistent spacing
- Vector auto-shape icons in circular badges
- Subtle fade transitions between slides

Run:  python3 build_love_presentation.py
Output: Love_The_Universal_Language.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


# ----------------------------------------------------------------------------
# Design tokens
# ----------------------------------------------------------------------------
class C:
    WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
    CREAM      = RGBColor(0xFF, 0xFB, 0xFC)   # near-white warm background
    BLUSH      = RGBColor(0xFB, 0xE4, 0xE8)   # light blush pink
    BLUSH_SOFT = RGBColor(0xFA, 0xEF, 0xF1)   # very soft blush
    ROSE       = RGBColor(0xD9, 0x5C, 0x77)   # light red / rose (primary accent)
    ROSE_DEEP  = RGBColor(0xB8, 0x40, 0x5C)   # deeper rose for emphasis
    GOLD       = RGBColor(0xC7, 0xA2, 0x4A)   # gold accent
    CHARCOAL   = RGBColor(0x3A, 0x2C, 0x30)   # primary text
    MUTE       = RGBColor(0x8C, 0x7B, 0x80)   # secondary / muted text
    LINE       = RGBColor(0xEB, 0xD9, 0xDD)   # hairline dividers


HEAD_FONT = "Georgia"
BODY_FONT = "Segoe UI"

EMU_IN = 914400
SLIDE_W = 13.333
SLIDE_H = 7.5


# ----------------------------------------------------------------------------
# Low-level helpers
# ----------------------------------------------------------------------------
def _set_alpha(shape, pct):
    """Apply opacity (0-100 where 100 = fully opaque) to a shape's solid fill."""
    spPr = shape._element.spPr
    solidFill = spPr.find(qn("a:solidFill"))
    if solidFill is None:
        return
    srgb = solidFill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for existing in srgb.findall(qn("a:alpha")):
        srgb.remove(existing)
    alpha = srgb.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))})
    srgb.append(alpha)


def add_rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE,
             line_color=None, line_w=None, alpha=None, shadow=False):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if alpha is not None:
        _set_alpha(sp, alpha)
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w or 1)
    sp.shadow.inherit = False
    if shadow:
        _soft_shadow(sp)
    return sp


def _soft_shadow(sp):
    """Add a subtle, designer-style outer shadow."""
    spPr = sp._element.spPr
    effLst = spPr.makeelement(qn("a:effectLst"), {})
    shdw = spPr.makeelement(qn("a:outerShdw"), {
        "blurRad": "90000", "dist": "38100", "dir": "5400000",
        "rotWithShape": "0",
    })
    clr = spPr.makeelement(qn("a:srgbClr"), {"val": "B8405C"})
    alpha = spPr.makeelement(qn("a:alpha"), {"val": "18000"})
    clr.append(alpha)
    shdw.append(clr)
    effLst.append(shdw)
    spPr.append(effLst)


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.0, space_after=0):
    """runs: list of paragraphs; each paragraph is a list of (text, dict) run specs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        for text, spec in para:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = spec.get("font", BODY_FONT)
            f.size = Pt(spec.get("size", 18))
            f.bold = spec.get("bold", False)
            f.italic = spec.get("italic", False)
            f.color.rgb = spec.get("color", C.CHARCOAL)
            if "spacing" in spec:
                _letter_spacing(r, spec["spacing"])
    return tb


def _letter_spacing(run, pts):
    """Set character spacing (tracking) in points."""
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(int(pts * 100)))


def _draw_home(slide, cx, cy, size, color):
    """Composite house icon (roof + body + door) centered at (cx, cy)."""
    w = size
    roof_h = size * 0.44
    body_h = size * 0.56
    # roof
    add_rect(slide, cx - w / 2, cy - size / 2, w, roof_h, color,
             shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    # body
    body_w = w * 0.78
    add_rect(slide, cx - body_w / 2, cy - size / 2 + roof_h, body_w, body_h, color,
             shape=MSO_SHAPE.RECTANGLE)
    # door (heart) - a warm family touch
    dsz = size * 0.24
    add_rect(slide, cx - dsz / 2, cy - size / 2 + roof_h + body_h - dsz * 1.05,
             dsz, dsz, C.WHITE, shape=MSO_SHAPE.HEART)


def badge(slide, cx, cy, d, fill, icon_shape, icon_color=C.WHITE,
          ring=None, icon_scale=0.5):
    """Circular badge centered at (cx, cy) with an auto-shape icon inside.

    icon_shape may be an MSO_SHAPE member or the string "HOME" for a
    composite house icon.
    """
    if ring is not None:
        rd = d + 0.18
        add_rect(slide, cx - rd / 2, cy - rd / 2, rd, rd, C.WHITE,
                 shape=MSO_SHAPE.OVAL, line_color=ring, line_w=1.25)
    circle = add_rect(slide, cx - d / 2, cy - d / 2, d, d, fill,
                      shape=MSO_SHAPE.OVAL, shadow=True)
    isz = d * icon_scale
    if icon_shape == "HOME":
        _draw_home(slide, cx, cy, isz, icon_color)
        return circle, None
    ic = add_rect(slide, cx - isz / 2, cy - isz / 2, isz, isz, icon_color,
                  shape=icon_shape)
    return circle, ic


def add_fade(slide, speed="med"):
    """Inject a fade transition into the slide XML."""
    sld = slide._element
    trans = sld.makeelement(qn("p:transition"), {"spd": speed})
    fade = sld.makeelement(qn("p:fade"), {})
    trans.append(fade)
    clrMapOvr = sld.find(qn("p:clrMapOvr"))
    ref = clrMapOvr if clrMapOvr is not None else sld.find(qn("p:cSld"))
    ref.addnext(trans)


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color=C.CREAM):
    add_rect(slide, -0.05, -0.05, SLIDE_W + 0.1, SLIDE_H + 0.1, color)


def deco_heart(slide, x, y, size, color, alpha):
    """A soft, semi-transparent decorative heart."""
    sp = add_rect(slide, x, y, size, size, color, shape=MSO_SHAPE.HEART, alpha=alpha)
    return sp


def kicker(slide, x, y, text, color=C.ROSE, align=PP_ALIGN.LEFT):
    """Small uppercase label with tracking - a designer touch."""
    add_text(slide, x, y, 6, 0.3, [[(text, {
        "font": BODY_FONT, "size": 12.5, "bold": True,
        "color": color, "spacing": 3.2})]], align=align)


def page_num(slide, n):
    add_text(slide, SLIDE_W - 1.2, SLIDE_H - 0.62, 0.8, 0.3,
             [[(f"{n:02d}", {"font": BODY_FONT, "size": 11, "color": C.MUTE,
                             "spacing": 1.5})]], align=PP_ALIGN.RIGHT)
    add_text(slide, 0.9, SLIDE_H - 0.62, 3.0, 0.3,
             [[("LOVE  \u2014  THE UNIVERSAL LANGUAGE", {
                 "font": BODY_FONT, "size": 9, "color": C.MUTE, "spacing": 1.8})]])


def gold_rule(slide, x, y, w=1.0):
    add_rect(slide, x, y, w, 0.035, C.GOLD)


# ----------------------------------------------------------------------------
# Slides
# ----------------------------------------------------------------------------
def slide_title(prs):
    s = blank_slide(prs)
    bg(s, C.CREAM)

    # Soft blush field on the right + large faint heart motif
    add_rect(s, 8.35, -0.5, 5.6, 8.5, C.BLUSH, alpha=55)
    deco_heart(s, 8.9, 1.15, 4.8, C.ROSE, alpha=10)
    deco_heart(s, 11.4, 4.7, 2.0, C.GOLD, alpha=12)
    deco_heart(s, 8.5, 5.6, 1.1, C.ROSE, alpha=14)

    # Small heart mark near the kicker
    add_rect(s, 0.92, 1.42, 0.34, 0.34, C.ROSE, shape=MSO_SHAPE.HEART)
    kicker(s, 1.42, 1.5, "A PRESENTATION ON HUMAN CONNECTION")

    gold_rule(s, 0.95, 2.28, 1.15)

    add_text(s, 0.9, 2.55, 8.6, 2.6, [
        [("Love:", {"font": HEAD_FONT, "size": 60, "bold": True, "color": C.CHARCOAL})],
        [("The Universal Language", {"font": HEAD_FONT, "size": 60, "bold": True,
                                     "color": C.ROSE})],
    ], line_spacing=1.02)

    add_text(s, 0.95, 5.15, 7.6, 0.9, [
        [("Understanding the Power of Human Connection", {
            "font": BODY_FONT, "size": 20, "color": C.MUTE})],
    ])

    page_num(s, 1)
    add_fade(s)
    return s


def content_header(slide, n, kick, title_parts, subtitle=None):
    """Shared header block for content slides."""
    add_rect(slide, 0.92, 0.86, 0.3, 0.3, C.ROSE, shape=MSO_SHAPE.HEART)
    kicker(slide, 1.36, 0.92, kick)
    add_text(slide, 0.9, 1.25, 11.5, 1.0, [title_parts],
             line_spacing=1.0)
    gold_rule(slide, 0.95, 2.16, 0.9)
    if subtitle:
        add_text(slide, 0.95, 2.32, 11.0, 0.5, [[(subtitle, {
            "font": BODY_FONT, "size": 15, "italic": True, "color": C.MUTE})]])


def slide_what_is_love(prs):
    s = blank_slide(prs)
    bg(s, C.CREAM)
    deco_heart(s, 11.6, 5.4, 2.6, C.BLUSH, alpha=45)

    content_header(
        s, 2, "SECTION 01",
        [("What Is Love?", {"font": HEAD_FONT, "size": 38, "bold": True,
                            "color": C.CHARCOAL})],
    )

    # Definition block
    add_text(s, 0.95, 2.55, 11.3, 0.95, [[
        ("Love is a deep and enduring bond of care, affection, and connection \u2014 ",
         {"font": BODY_FONT, "size": 18, "color": C.CHARCOAL}),
        ("a choice to value another\u2019s well-being as our own.",
         {"font": BODY_FONT, "size": 18, "italic": True, "color": C.ROSE_DEEP}),
    ]], line_spacing=1.15)

    # Four type cards
    types = [
        ("Romantic Love", "Passion, intimacy, and\ndevoted partnership.", MSO_SHAPE.HEART, C.ROSE),
        ("Family Love", "Unconditional bonds\nthat shape who we are.", "HOME", C.ROSE_DEEP),
        ("Friendship", "Loyalty, trust, and\nshared joy over time.", MSO_SHAPE.SMILEY_FACE, C.GOLD),
        ("Self-Love", "Compassion and respect\nfor one\u2019s own worth.", MSO_SHAPE.SUN, C.ROSE),
    ]
    n = len(types)
    card_w, gap = 2.72, 0.28
    total = n * card_w + (n - 1) * gap
    x0 = (SLIDE_W - total) / 2
    top, card_h = 3.75, 2.85
    for i, (name, desc, icon, col) in enumerate(types):
        x = x0 + i * (card_w + gap)
        add_rect(s, x, top, card_w, card_h, C.WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 line_color=C.LINE, line_w=1, shadow=True)
        badge(s, x + card_w / 2, top + 0.82, 0.98, col, icon, icon_scale=0.46)
        add_text(s, x + 0.15, top + 1.42, card_w - 0.3, 0.4, [[(name, {
            "font": HEAD_FONT, "size": 16.5, "bold": True, "color": C.CHARCOAL})]],
            align=PP_ALIGN.CENTER)
        add_text(s, x + 0.2, top + 1.92, card_w - 0.4, 0.8, [[(desc, {
            "font": BODY_FONT, "size": 12.5, "color": C.MUTE})]],
            align=PP_ALIGN.CENTER, line_spacing=1.1)

    page_num(s, 2)
    add_fade(s)
    return s


def slide_why_matters(prs):
    s = blank_slide(prs)
    bg(s, C.CREAM)
    add_rect(s, 8.6, -0.5, 5.5, 8.5, C.BLUSH_SOFT, alpha=70)
    deco_heart(s, 10.9, 0.7, 2.2, C.ROSE, alpha=10)

    content_header(
        s, 3, "SECTION 02",
        [("Why Love Matters", {"font": HEAD_FONT, "size": 38, "bold": True,
                               "color": C.CHARCOAL})],
        subtitle="Connection is one of the strongest predictors of a healthy, meaningful life.",
    )

    points = [
        ("Mental Health", "Lowers stress and anxiety while boosting happiness and resilience.",
         MSO_SHAPE.SMILEY_FACE, C.ROSE),
        ("Physical Well-Being", "Strong bonds support heart health, immunity, and longevity.",
         MSO_SHAPE.HEART, C.ROSE_DEEP),
        ("Relationships", "Builds trust, belonging, and communities that thrive together.",
         MSO_SHAPE.STAR_5_POINT, C.GOLD),
        ("Personal Growth", "Encourages empathy, purpose, and the courage to evolve.",
         MSO_SHAPE.UP_ARROW, C.ROSE),
    ]
    top = 3.05
    row_h = 1.02
    for i, (title, desc, icon, col) in enumerate(points):
        y = top + i * row_h
        badge(s, 1.35, y + 0.36, 0.72, col, icon, icon_scale=0.5)
        add_text(s, 2.05, y - 0.04, 9.9, 0.45, [[(title, {
            "font": HEAD_FONT, "size": 17, "bold": True, "color": C.CHARCOAL})]])
        add_text(s, 2.05, y + 0.36, 9.7, 0.5, [[(desc, {
            "font": BODY_FONT, "size": 13.5, "color": C.MUTE})]])
        if i < len(points) - 1:
            add_rect(s, 2.05, y + 0.9, 9.5, 0.014, C.LINE)

    page_num(s, 3)
    add_fade(s)
    return s


def slide_express(prs):
    s = blank_slide(prs)
    bg(s, C.CREAM)
    deco_heart(s, 0.2, 5.2, 2.4, C.BLUSH, alpha=45)

    content_header(
        s, 4, "SECTION 03",
        [("Ways to Express Love", {"font": HEAD_FONT, "size": 38, "bold": True,
                                   "color": C.CHARCOAL})],
        subtitle="Small, consistent acts communicate love more powerfully than words alone.",
    )

    cards = [
        ("Kind Words", "Encourage and\naffirm sincerely.", MSO_SHAPE.HEART, C.ROSE),
        ("Quality Time", "Be fully present,\nwithout distraction.", MSO_SHAPE.SUN, C.GOLD),
        ("Acts of Kindness", "Help freely, expecting\nnothing in return.", MSO_SHAPE.HEART, C.ROSE_DEEP),
        ("Listening", "Seek to understand\nbefore responding.", MSO_SHAPE.SMILEY_FACE, C.ROSE),
        ("Respect & Support", "Honor others and\nstand beside them.", MSO_SHAPE.STAR_5_POINT, C.GOLD),
    ]
    n = len(cards)
    card_w, gap = 2.24, 0.22
    total = n * card_w + (n - 1) * gap
    x0 = (SLIDE_W - total) / 2
    top, card_h = 3.15, 3.35
    for i, (name, desc, icon, col) in enumerate(cards):
        x = x0 + i * (card_w + gap)
        add_rect(s, x, top, card_w, card_h, C.WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 line_color=C.LINE, line_w=1, shadow=True)
        # accent top bar
        add_rect(s, x, top, card_w, 0.12, col, shape=MSO_SHAPE.ROUND_1_RECTANGLE)
        # step number
        add_text(s, x + 0.22, top + 0.28, card_w - 0.44, 0.4, [[(f"0{i+1}", {
            "font": HEAD_FONT, "size": 15, "bold": True, "color": col, "spacing": 1})]])
        badge(s, x + card_w / 2, top + 1.25, 0.86, col, icon, icon_scale=0.46)
        add_text(s, x + 0.12, top + 1.85, card_w - 0.24, 0.4, [[(name, {
            "font": HEAD_FONT, "size": 14.5, "bold": True, "color": C.CHARCOAL})]],
            align=PP_ALIGN.CENTER)
        add_text(s, x + 0.16, top + 2.35, card_w - 0.32, 0.8, [[(desc, {
            "font": BODY_FONT, "size": 11.5, "color": C.MUTE})]],
            align=PP_ALIGN.CENTER, line_spacing=1.1)

    page_num(s, 4)
    add_fade(s)
    return s


def slide_conclusion(prs):
    s = blank_slide(prs)
    bg(s, C.CREAM)
    # Elegant blush panel with layered hearts
    add_rect(s, -0.05, -0.05, SLIDE_W + 0.1, SLIDE_H + 0.1, C.BLUSH, alpha=40)
    deco_heart(s, 10.4, 0.5, 3.4, C.ROSE, alpha=10)
    deco_heart(s, 0.4, 4.4, 2.8, C.GOLD, alpha=10)

    add_rect(s, SLIDE_W / 2 - 0.2, 0.85, 0.4, 0.4, C.ROSE, shape=MSO_SHAPE.HEART)
    add_text(s, 0, 1.4, SLIDE_W, 0.3, [[("IN CLOSING", {
        "font": BODY_FONT, "size": 12.5, "bold": True,
        "color": C.ROSE, "spacing": 3.2})]], align=PP_ALIGN.CENTER)

    add_text(s, 1.5, 1.95, SLIDE_W - 3.0, 1.4, [[
        ("Love is not only a feeling \u2014 it is a ", {
            "font": HEAD_FONT, "size": 27, "color": C.CHARCOAL}),
        ("choice", {"font": HEAD_FONT, "size": 27, "bold": True, "italic": True,
                    "color": C.ROSE}),
        (" expressed through care, respect, and compassion.", {
            "font": HEAD_FONT, "size": 27, "color": C.CHARCOAL}),
    ]], align=PP_ALIGN.CENTER, line_spacing=1.12)

    # Divider with gold accent centered
    add_rect(s, SLIDE_W / 2 - 0.6, 3.75, 1.2, 0.04, C.GOLD)

    # Inspirational quote
    add_text(s, 2.0, 4.05, SLIDE_W - 4.0, 1.0, [
        [("\u201CWhere there is love, there is life.\u201D", {
            "font": HEAD_FONT, "size": 23, "italic": True, "color": C.ROSE_DEEP})],
        [("\u2014 Mahatma Gandhi", {
            "font": BODY_FONT, "size": 14, "color": C.MUTE, "spacing": 1.5})],
    ], align=PP_ALIGN.CENTER, line_spacing=1.2, space_after=6)

    # Thank you section
    add_text(s, 0, 5.55, SLIDE_W, 1.0, [[("Thank You", {
        "font": HEAD_FONT, "size": 40, "bold": True, "color": C.CHARCOAL})]],
        align=PP_ALIGN.CENTER)
    add_text(s, 0, 6.55, SLIDE_W, 0.4, [[(
        "Choose love, and share it generously.", {
            "font": BODY_FONT, "size": 14, "color": C.MUTE, "spacing": 0.8})]],
        align=PP_ALIGN.CENTER)

    add_fade(s)
    return s


# ----------------------------------------------------------------------------
# Build
# ----------------------------------------------------------------------------
def build():
    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W * EMU_IN))
    prs.slide_height = Emu(int(SLIDE_H * EMU_IN))

    slide_title(prs)
    slide_what_is_love(prs)
    slide_why_matters(prs)
    slide_express(prs)
    slide_conclusion(prs)

    # Core document properties
    cp = prs.core_properties
    cp.title = "Love: The Universal Language"
    cp.subject = "Understanding the Power of Human Connection"
    cp.author = "Kiro"
    cp.category = "Presentation"
    cp.comments = "Understanding the Power of Human Connection"

    out = "Love_The_Universal_Language.pptx"
    prs.save(out)
    print(f"Saved {out} with {len(prs.slides._sldIdLst)} slides.")


if __name__ == "__main__":
    build()
